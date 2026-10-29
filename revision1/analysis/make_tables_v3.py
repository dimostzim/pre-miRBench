#!/usr/bin/env python3
"""
v3 revision-1 tables for co-authors (make_tables.py of v2 with the v3 methods, labels and notes): every predictor, retrained
MirMachine and the shortcut baselines on the four benchmark tests, and every method genome-wide. One model per method.

Inputs (--results-dir, the layout of node 4 v3/results):
  benchmark/summary/summary.csv   summarize_benchmark.py output (whole-split rows)
  genome/metrics/<tool>_<sp>.csv  genome_scale_metrics.py output (one model per method)
  genome/comparators/*.summary.csv compare_annotation_calls.py output (homology methods, their own calls)
  benchmark/summary/replicate_metrics.csv, paired_differences.csv, benchmark/validation_ap.csv
  genome/family_ap/<tool>_<sp>.csv  genome_family_ap.py output (AP for training-family and held-out-family loci)
Outputs: --output-dir/benchmark_tests.csv, genome_wide.csv (figure inputs), benchmark_tests_all_metrics.tsv (every
summary metric, one row per method and test), genome_wide_all_metrics.tsv (every genome metric, one row per method,
genome and calling rule), all_results.tsv (every number, long format) and --xlsx.
A genome result that is not yet available is written as "pending".
"""
import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TESTS = [
    ("test_known_species_known_family", "Test 1: known species, known family"),
    ("test_known_species_heldout_family", "Test 2: known species, held-out family"),
    ("test_heldout_species_known_family", "Test 3: held-out species, known family"),
    ("test_heldout_species_heldout_family", "Test 4: held-out species, held-out family"),
]
GENOMES = [("dme", "D. melanogaster"), ("dan", "D. ananassae"), ("dya", "D. yakuba"), ("gga", "G. gallus (chicken)")]
PREDICTORS = {  # genome-wide rows follow this order: the benchmark sheet's order (Test 4 AP, highest first)
    "premirbench_model_v3": "pre-miRBench model v3 (Agentomics)",
    "mustard": "MuStARD",
    "mirdnn": "mirDNN",
    "deepmirgene": "deepMiRGene",
    "dnnpremir": "dnnPreMiR",
    "mire2e": "miRe2e",
    "deepmir": "DeepMir",
}
HOMOLOGY = {
    "mirmachine_retrained_genome_cutoffs": "MirMachine, retrained (cutoffs fitted on training genomes)",
    "ncortho_training_family_refs": "ncOrtho, training-family references",
}
REFERENCE = {}  # off-the-shelf MirMachine and all-reference ncOrtho were not run on v3
NOT_REPORTED = {"premirbench_model_v2"}  # the v2 model recipe retrained on v3: in the summary files, left out of the tables (user, 25 Oct)
BASELINES = {
    "baseline_combined_gb": "Mining features + positional bins + loop distance from centre (gradient boosting)",
    "baseline_paired_profile_lr": "Paired fraction in 20 positional bins (logistic regression)",
    "baseline_mining_rf": "22 mining features (random forest)",
    "baseline_kmer_lr": "k-mer composition, k = 1-4 (logistic regression)",
    "baseline_hairpin_filter_lr": "Four hairpin-filter features (logistic regression)",
    "baseline_mfe": "MFE only (score = -MFE, no fitting)",
    "baseline_loop_centrality_lr": "Loop distance from centre (logistic regression)",
}
GENOME_MODEL = {tool: "seed 1" for tool in PREDICTORS} | {"premirbench_model_v3": "agent's networks"}
GENOME_MODEL |= {tool: "single run" for tool in [*HOMOLOGY, *REFERENCE]}
COMPARATOR_FILES = {("mirmachine_retrained_genome_cutoffs", sp): f"mirmachine_retrained_genome_cutoffs_{sp}.summary.csv"
                    for sp, _ in GENOMES}
COMPARATOR_FILES[("ncortho_training_family_refs", "gga")] = "ncortho_training_family_refs_gga.summary.csv"
NAMES = {**PREDICTORS, **HOMOLOGY, **BASELINES, **REFERENCE}
CATEGORY = {**{t: "predictor" for t in PREDICTORS}, **{t: "homology" for t in HOMOLOGY},
            **{t: "shortcut baseline" for t in BASELINES}, **{t: "reference, not comparable" for t in REFERENCE}}
FAMILY_FIELDS = ["ap_training_family_loci", "ap_heldout_family_loci", "positives_training_family", "positives_heldout_family"]
GENOME_FIELDS = ["calls", "calls_per_mb", "precision_unambiguous_calls", "ap_unambiguous_candidates", *FAMILY_FIELDS, "all_loci_recall",
                 "unmasked_loci_recall", "unmasked_training_family_loci_recall", "unmasked_heldout_family_loci_recall",
                 "all_loci", "unmasked_loci", "unmasked_training_family_loci", "unmasked_heldout_family_loci"]


def read_csv(path):
    with open(path, newline="") as handle:
        return list(csv.DictReader(handle))


def family_ap(results, tool, sp, ap_all):
    """AP by family group (genome_family_ap.py); checks its all-loci AP against the genome tables."""
    path = results / f"genome/family_ap/{tool}_{sp}.csv"
    if not path.exists():
        return {field: "pending" for field in FAMILY_FIELDS}
    row = read_csv(path)[0]
    if ap_all not in ("", "pending") and abs(float(row["ap_unambiguous_candidates"]) - float(ap_all)) > 1e-9:
        raise ValueError(f"{path}: AP {row['ap_unambiguous_candidates']} differs from the genome tables ({ap_all})")
    return {field: row[field] for field in FAMILY_FIELDS}


def with_family(row, family):
    """Insert the family-group fields right after the all-loci AP."""
    out = {}
    for key, value in row.items():
        out[key] = value
        if key == "ap_unambiguous_candidates":
            out.update(family)
    return out


def benchmark_rows(results):
    summary = {(r["tool"], r["split"]): r for r in read_csv(results / "benchmark/summary/summary.csv") if r["group"] == ""}
    rows = []
    benchmark_homology = {tool: HOMOLOGY[tool] for tool in ("mirmachine_retrained", "mirmachine_retrained_genome_cutoffs")
                          if (tool, TESTS[0][0]) in summary}
    for category, names in [("predictor", PREDICTORS), ("homology", benchmark_homology), ("shortcut baseline", BASELINES)]:
        for tool, name in names.items():
            for split, _ in TESTS:
                s = summary[(tool, split)]
                rows.append({"category": category, "tool": tool, "method": name, "test": split, "replicates": s["replicates"],
                             "positives": s["positives"], "negatives": s["negatives"],
                             "ap_mean": s["ap_mean"], "ap_sd": s["ap_sd"], "ap_ci_low": s["ap_ci_low"], "ap_ci_high": s["ap_ci_high"],
                             "auroc_mean": s["auroc_mean"],
                             "valid_max_f1_precision_mean": s["valid_max_f1_precision_mean"],
                             "valid_max_f1_recall_mean": s["valid_max_f1_recall_mean"],
                             "valid_max_f1_f1_mean": s["valid_max_f1_f1_mean"]})
    return rows


def genome_rows(results, groups=None):
    rows = []
    for category, names in groups or [("predictor", PREDICTORS), ("homology", HOMOLOGY)]:
        for tool, name in names.items():
            for sp, _ in GENOMES:
                row = {"category": category, "tool": tool, "method": name, "genome": sp, "model": GENOME_MODEL[tool]}
                if category == "predictor":
                    path = results / f"genome/metrics/{tool}_{sp}.csv"
                    row["calls_by"] = "threshold: validation max F1"
                    found = [r for r in read_csv(path) if r["threshold_rule"] == "valid_max_f1"] if path.exists() else []
                else:
                    if (tool, sp) not in COMPARATOR_FILES:
                        continue  # ncOrtho was run on chicken only
                    path = results / "genome/comparators" / COMPARATOR_FILES[(tool, sp)]
                    row["calls_by"] = "own calls"
                    found = read_csv(path) if path.exists() else []
                if len(found) > 1:
                    raise ValueError(f"{path}: expected one row, found {len(found)}")
                for field in GENOME_FIELDS:
                    if field not in FAMILY_FIELDS:
                        row[field] = found[0][field] if found else "pending"
                if found:
                    row.update(family_ap(results, tool, sp, row["ap_unambiguous_candidates"]))
                else:
                    row.update({field: "pending" for field in FAMILY_FIELDS})
                rows.append({field: row[field] for field in [*row][:6] + GENOME_FIELDS})
    return rows


def validation_rows(results):
    return [{"category": CATEGORY[r["tool"]], "tool": r["tool"], "method": NAMES[r["tool"]], "model": r["replicate"],
             "positives": r["positives"], "negatives": r["negatives"], "ap": r["ap"], "auroc": r["auroc"]}
            for r in read_csv(results / "benchmark/validation_ap.csv") if r["tool"] not in NOT_REPORTED]


def replicate_rows(results):
    """Every replicate and test; species = "" for the whole test, else one held-out species."""
    return [{"category": CATEGORY[r["tool"]], "tool": r["tool"], "method": NAMES[r["tool"]], "model": r["replicate"], "test": r["split"],
             "species": r["group"], **{k: v for k, v in r.items() if k not in ("tool", "replicate", "group", "split")}}
            for r in read_csv(results / "benchmark/summary/replicate_metrics.csv") if r["tool"] not in NOT_REPORTED]


def summary_rows(results):
    """Every summary.csv row (mean, SD and bootstrap CI over replicates); species = "" for the whole test."""
    return [{"category": CATEGORY[r["tool"]], "tool": r["tool"], "method": NAMES[r["tool"]], "test": r["split"], "species": r["group"],
             **{k: v for k, v in r.items() if k not in ("tool", "group", "split")}}
            for r in read_csv(results / "benchmark/summary/summary.csv") if r["tool"] not in NOT_REPORTED]


def auto_columns(rows, skip=("tool", "tool_a", "tool_b")):
    """Union of the rows' fields in order of first appearance."""
    fields = list(dict.fromkeys(field for row in rows for field in row))
    return [(field, field.replace("_", " ")) for field in fields if field not in skip]


def species_rows(results):
    return [{"category": CATEGORY[r["tool"]], "tool": r["tool"], "method": NAMES[r["tool"]], "test": r["split"], "species": r["group"],
             "replicates": r["replicates"], "positives": r["positives"], "negatives": r["negatives"],
             "ap_mean": r["ap_mean"], "ap_sd": r["ap_sd"], "auroc_mean": r["auroc_mean"]}
            for r in read_csv(results / "benchmark/summary/summary.csv") if r["group"] != "" and r["tool"] not in NOT_REPORTED]


def paired_rows(results):
    return [{"test": r["split"], "method_a": NAMES[r["tool_a"]], "method_b": NAMES[r["tool_b"]], "tool_a": r["tool_a"], "tool_b": r["tool_b"],
             "replicates_a": r["replicates_a"], "replicates_b": r["replicates_b"], "ap_difference_a_minus_b": r["ap_difference"],
             "ci_low": r["ci_low"], "ci_high": r["ci_high"]}
            for r in read_csv(results / "benchmark/summary/paired_differences.csv")
            if r["tool_a"] not in NOT_REPORTED and r["tool_b"] not in NOT_REPORTED]


def genome_all_rows(results):
    """Every genome metric: predictors at both validation thresholds, homology methods and reference methods."""
    rows = []
    for tool in PREDICTORS:
        for sp, _ in GENOMES:
            path = results / f"genome/metrics/{tool}_{sp}.csv"
            for r in (read_csv(path) if path.exists() else []):
                rows.append(with_family({"category": CATEGORY[tool], "tool": tool, "method": NAMES[tool], "genome": sp, "model": GENOME_MODEL[tool],
                                         "calls_by": f"threshold: {r['threshold_rule']}", **{k: v for k, v in r.items() if k not in ("tool", "threshold_rule")}},
                                        family_ap(results, tool, sp, r["ap_unambiguous_candidates"])))
    for tool in [*HOMOLOGY, *REFERENCE]:
        for sp, _ in GENOMES:
            if (tool, sp) not in COMPARATOR_FILES:
                continue
            path = results / "genome/comparators" / COMPARATOR_FILES[(tool, sp)]
            for r in (read_csv(path) if path.exists() else []):
                rows.append(with_family({"category": CATEGORY[tool], "tool": tool, "method": NAMES[tool], "genome": sp, "model": GENOME_MODEL[tool],
                                         "calls_by": "own calls", **{k: v for k, v in r.items() if k != "tool"}},
                                        family_ap(results, tool, sp, r["ap_unambiguous_candidates"])))
    return rows


def long_rows(sections):
    """One row per number: section, category, method, model, dataset, group, metric, value."""
    keys = {"category", "tool", "method", "model", "test", "genome", "species", "calls_by", "method_a", "method_b", "tool_a", "tool_b",
            "replicates", "replicates_a", "replicates_b", "split", "threshold"}
    out = []
    for section, rows in sections:
        for r in rows:
            method_id = r.get("tool") or f"{r['tool_a']} vs {r['tool_b']}"
            method = r.get("method") or f"{r['method_a']} minus {r['method_b']}"
            dataset = r.get("test") or r.get("genome") or ("valid" if section == "benchmark_validation" else "")
            group = r.get("species") or r.get("calls_by") or ""
            for metric, value in r.items():
                if metric in keys or value in ("", "pending"):
                    continue
                out.append({"section": section, "category": r.get("category", "pair"), "method_id": method_id, "method": method,
                            "model": r.get("model", "one model" if section in ("benchmark_tests", "benchmark_per_species") else ""),
                            "dataset": dataset, "group": group, "metric": metric, "value": value})
    return out


def write_csv(path, rows, delimiter=","):
    with open(path, "w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), delimiter=delimiter)
        writer.writeheader()
        writer.writerows(rows)


FONT = "Arial"
BOLD = Font(name=FONT, bold=True)
NORMAL = Font(name=FONT)
TITLE = Font(name=FONT, bold=True, size=12)
ITALIC = Font(name=FONT, italic=True, color="555555")
HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
SECTION_FILL = PatternFill("solid", fgColor="F2F2F2")
OWN_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Border(bottom=Side(style="thin", color="999999"))
WRAP = Alignment(wrap_text=True, vertical="top")


def fmt(value, digits=3):
    return f"{float(value):.{digits}f}"


def put(ws, row, column, value, font=NORMAL, fill=None, number_format=None, wrap=False):
    cell = ws.cell(row=row, column=column, value=value)
    cell.font = font
    if fill:
        cell.fill = fill
    if number_format:
        cell.number_format = number_format
    if wrap:
        cell.alignment = WRAP
    return cell


def header(ws, row, labels):
    for column, label in enumerate(labels, 1):
        cell = put(ws, row, column, label, BOLD, HEADER_FILL, wrap=True)
        cell.border = THIN


def benchmark_sheet(wb, rows):
    ws = wb.active
    ws.title = "Benchmark tests"
    put(ws, 1, 1, "pre-miRBench v3 benchmark: AP (average precision) on the four test sets", TITLE)
    put(ws, 2, 1, "One model per method: AP [95% bootstrap CI]. Positives : negatives = 1 : 10 in every test, "
                  "so chance-level AP is about 0.09. Highlighted row: our method.", ITALIC)
    header(ws, 4, ["Method", "Replicates"] + [label for _, label in TESTS])
    by = {(r["tool"], r["test"]): r for r in rows}
    first = rows[0]
    put(ws, 5, 1, "Records (positives / negatives)", ITALIC)
    for column, (split, _) in enumerate(TESTS, 3):
        r = by[(first["tool"], split)]
        put(ws, 5, column, f"{r['positives']} / {r['negatives']}", ITALIC)
    line = 6
    for category, title in [("predictor", "Predictors, retrained from scratch on the v3 training partition"),
                            ("homology", "Homology-based annotation"),
                            ("shortcut baseline", "Shortcut baselines (reference only; simple features that reveal benchmark cues)")]:
        for column in range(1, 7):
            put(ws, line, column, title if column == 1 else None, BOLD, SECTION_FILL)
        line += 1
        tools = list(dict.fromkeys(r["tool"] for r in rows if r["category"] == category))
        tools.sort(key=lambda t: -float(by[(t, TESTS[3][0])]["ap_mean"]))
        for tool in tools:
            fill = OWN_FILL if tool == "premirbench_model_v3" else None
            put(ws, line, 1, by[(tool, TESTS[0][0])]["method"], BOLD if fill else NORMAL, fill, wrap=True)
            put(ws, line, 2, int(by[(tool, TESTS[0][0])]["replicates"]), NORMAL, fill)
            for column, (split, _) in enumerate(TESTS, 3):
                r = by[(tool, split)]
                sd = "" if r["ap_sd"] in ("", "nan") else f" ± {fmt(r['ap_sd'])}"
                text = f"{fmt(r['ap_mean'])}{sd} [{fmt(r['ap_ci_low'])}–{fmt(r['ap_ci_high'])}]"
                put(ws, line, column, text, BOLD if fill else NORMAL, fill)
            line += 1
    widths = [62, 11, 30, 30, 30, 30]
    for column, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "C5"


def numeric_sheet(wb, title, rows, columns):
    ws = wb.create_sheet(title)
    header(ws, 1, [label for _, label in columns])
    for line, row in enumerate(rows, 2):
        for column, (field, _) in enumerate(columns, 1):
            value = row.get(field, "")
            try:
                value = int(value) if field in {"replicates", "positives", "negatives", "calls", "all_loci", "unmasked_loci",
                                                 "unmasked_training_family_loci", "unmasked_heldout_family_loci"} else float(value)
            except ValueError:
                pass  # text fields and "pending"
            put(ws, line, column, value, number_format="0.0000" if isinstance(value, float) else None)
    for column, (field, label) in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(column)].width = 58 if field == "method" else max(12, min(24, len(label) + 2))
    ws.freeze_panes = "B2"


def genome_sheet(wb, rows):
    ws = wb.create_sheet("Genome-wide")
    put(ws, 1, 1, "Genome-wide evaluation on the four held-out genomes (labels: MirGeneDB loci only)", TITLE)
    put(ws, 2, 1, "Predictors: one model each (pre-miRBench model v3: the agent's networks), calls at the validation max-F1 threshold. "
                  "Homology methods: their own calls. Highlighted row: our method.", ITALIC)
    blocks = [("ap_unambiguous_candidates", "AP over unambiguous candidates (genomic prevalence)", "{:.3f}"),
              ("ap_training_family_loci", "AP, loci of training families (held-out-family loci left out)", "{:.3f}"),
              ("ap_heldout_family_loci", "AP, loci of held-out families (training-family loci left out)", "{:.3f}"),
              ("calls_per_mb", "Calls per Mb", None),
              ("precision_unambiguous_calls", "Precision of calls (unambiguous candidates)", "{:.4f}"),
              ("unmasked_loci_recall", "Recall of unmasked MirGeneDB loci", "{:.3f}"),
              ("unmasked_training_family_loci_recall", "Recall of unmasked loci, training families", "{:.3f}"),
              ("unmasked_heldout_family_loci_recall", "Recall of unmasked loci, held-out families", "{:.3f}")]
    by = {(r["tool"], r["genome"]): r for r in rows}
    tools = list(dict.fromkeys(r["tool"] for r in rows))
    line = 4
    for field, title, pattern in blocks:
        header(ws, line, [title, "Model", "Calls by"] + [label for _, label in GENOMES])
        line += 1
        for tool in tools:
            fill = OWN_FILL if tool == "premirbench_model_v3" else None
            sample = next(r for r in rows if r["tool"] == tool)
            put(ws, line, 1, sample["method"], BOLD if fill else NORMAL, fill)
            put(ws, line, 2, sample["model"], NORMAL, fill)
            put(ws, line, 3, sample["calls_by"], NORMAL, fill)
            for column, (sp, _) in enumerate(GENOMES, 4):
                r = by.get((tool, sp))
                if r is None:
                    text = "not run"
                elif r[field] == "pending":
                    text = "pending"
                else:
                    value = float(r[field])
                    text = (pattern or ("{:.1f}" if value >= 1 else "{:.2f}")).format(value)
                put(ws, line, column, text, BOLD if fill else NORMAL, fill)
            line += 1
        line += 1
    loci = next(r for r in rows if r["tool"] == "mirdnn")
    put(ws, line, 1, "MirGeneDB loci per genome (all / unmasked / unmasked training families / unmasked held-out families)", ITALIC)
    for column, (sp, _) in enumerate(GENOMES, 4):
        r = by[("mirdnn", sp)]
        put(ws, line, column, f"{r['all_loci']} / {r['unmasked_loci']} / {r['unmasked_training_family_loci']} / "
                              f"{r['unmasked_heldout_family_loci']}", ITALIC)
    for column, width in enumerate([58, 13, 26, 18, 18, 18, 20], 1):
        ws.column_dimensions[get_column_letter(column)].width = width


NOTES = [
    ("Benchmark", "v3 benchmark (revision 1): positive and negative windows follow one rule: 200 nt centred on the terminal loop "
                  "of a predicted hairpin (for a positive, the hairpin of its MirGeneDB precursor whose loop centre lies within "
                  "15 nt of the precursor midpoint), then shifted by a seeded 0-10 nt. Hairpins with more than 10% masked bases, "
                  "and windows with an N, are left out. Similarity filter against MirGeneDB precursors, no positive overlap "
                  "across partitions, four held-out species (gga, dme, dan, dya). Main metric: AP (average precision); ties are "
                  "one threshold, as in scikit-learn."),
    ("Predictors", "The six published predictors were retrained from scratch on the v3 training partition with their published "
                   "architectures and hyperparameters, one model each (seed 1), on the 4 July 2026 Docker images. DeepMir, "
                   "deepMiRGene, dnnPreMiR and mirDNN: our training loop around the authors' model builders, early stopping on "
                   "validation AUPRC (patience 20, at most 600 epochs). miRe2e and MuStARD: their own training loops, with "
                   "checkpoint selection on validation AUPRC (MuStARD's validation scoring runs in batches of 1,024 records to "
                   "avoid a GPU memory leak; validation scores equal within 5e-7). Inference on CPU."),
    ("pre-miRBench model v3", "Model selected by an Agentomics run on the v3 partitions (20 iterations; iteration 19 chosen on "
                              "the validation partition): the mean of three convolutional networks. The agent's trained "
                              "networks are used (validation AP 0.9740). Retraining with the released train.py on the "
                              "agent's GPU type (A10) rebuilds two of the three networks exactly; the model retrained that way "
                              "has validation AP 0.9722."),
    ("Retrained MirMachine", "MirMachine's covariance-model recipe applied to the v3 training partition only (711 family "
                             "models). Trusted cutoffs fitted on genome-wide searches of the training species as MirMachine "
                             "fitted its own cutoffs, with every hit near a validation or test window left out; 204 families "
                             "without a fitting hit keep the cutoff fitted on the training windows on the benchmark. "
                             "Benchmark score = best bitscore minus the family cutoff; one deterministic run, so no SD. It has "
                             "no model for held-out families, so it makes no calls on Tests 2 and 4 by construction (AP near "
                             "chance)."),
    ("ncOrtho", "Chicken genome only (human reference, four primate core species, ncOrtho 0.4.7; the run of 19 October). Only "
                "predictions from human references whose family is a v3 training family are counted (189 of 202 predictions), "
                "so held-out families are excluded as for the predictors."),
    ("Genome-wide", "Candidates: windows centred on every predicted hairpin of the genome, annotated loci not excluded "
                    "(1.1-6.5 million per genome; prevalence about 1:8,500 to 1:32,000). Positive: same strand as a MirGeneDB "
                    "precursor, with the window centre near the precursor midpoint; ambiguous: any other candidate overlapping a "
                    "precursor, excluded from AP and precision but counted as a call. AP over unambiguous (positive and negative) "
                    "candidates; homology methods rank candidates by the best overlapping bitscore, candidates without a hit "
                    "last. AP by family group: the same AP with all negatives and only the positives of training families "
                    "(held-out-family positives left out), or only those of held-out families; a positive's family is the "
                    "family of its MirGeneDB precursor. Chance-level AP is the share of positives among the candidates, below "
                    "0.0002 in every genome. Predictor calls: score at or above the threshold that maximises F1 "
                    "on the validation partition (1:10 prevalence). Precision: called unambiguous candidates that are MirGeneDB "
                    "positives. Recall: fraction of MirGeneDB loci found (predictors: the locus's candidate is called; homology "
                    "methods: a call overlaps the locus on the same strand). Unmasked loci: precursors with at most 10% masked "
                    "bases and no N. Training / held-out families: v3 training partition families."),
    ("Bootstrap", "10,000 class-stratified resamples of each test (seed 42), the same records for every model; the interval is "
                  "conditional on the trained model (one model per method, so training variation is not included)."),
    ("Not equalised", "The pre-miRBench model's architecture was selected by an agent on the validation partition; the published "
                      "predictors keep their published architectures and hyperparameters. One model per method."),
    ("Other sheets", "Validation AP: every trained model on the validation partition. Per model: every metric of every "
                     "model on each test, at both validation thresholds and at 0.5. Per species: the two held-out-species "
                     "tests split by species. Paired differences: AP of method A minus method B with its 95% "
                     "bootstrap interval. Genome-wide (all metrics): both validation thresholds and top-N recall. all_results.tsv "
                     "holds every number in long format (section, method, model, dataset, group, metric, value)."),
    ("Source", "Node biogemt-4, /SCRATCH/dtzim01/premirbench/v3/results (README.md, manifest.tsv with the md5 and origin of every "
               "input file). Generated by results/scripts/make_tables_v3.py."),
]


def notes_sheet(wb):
    ws = wb.create_sheet("Notes")
    header(ws, 1, ["Topic", "Note"])
    for line, (topic, text) in enumerate(NOTES, 2):
        put(ws, line, 1, topic, BOLD, wrap=True)
        put(ws, line, 2, text, wrap=True)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 120


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--results-dir", required=True)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--xlsx", required=True)
    args = parser.parse_args()
    results, out = Path(args.results_dir), Path(args.output_dir)
    out.mkdir(parents=True, exist_ok=True)
    bench, genome = benchmark_rows(results), genome_rows(results)
    write_csv(out / "benchmark_tests.csv", bench)
    write_csv(out / "genome_wide.csv", genome)
    labels = dict(TESTS)
    write_csv(out / "benchmark_tests_all_metrics.tsv",
              [{"category": r["category"], "method": r["method"], "tool": r["tool"], "test": r["test"], "test_name": labels[r["test"]],
                **{k: v for k, v in r.items() if k not in ("category", "method", "tool", "test", "species")}}
               for r in summary_rows(results) if r["species"] == ""], delimiter="\t")
    genome_all_with_reference = genome_all_rows(results)
    columns = list(dict.fromkeys(field for row in genome_all_with_reference for field in row))
    write_csv(out / "genome_wide_all_metrics.tsv", [{field: row.get(field, "") for field in columns} for row in genome_all_with_reference],
              delimiter="\t")

    wb = Workbook()
    benchmark_sheet(wb, bench)
    genome_sheet(wb, genome)
    numeric_sheet(wb, "Benchmark (numbers)", bench, [
        ("method", "Method"), ("category", "Category"), ("test", "Test"), ("replicates", "Replicates"),
        ("positives", "Positives"), ("negatives", "Negatives"), ("ap_mean", "AP, mean"), ("ap_sd", "AP, SD"),
        ("ap_ci_low", "AP, 95% CI low"), ("ap_ci_high", "AP, 95% CI high"), ("auroc_mean", "AUROC, mean"),
        ("valid_max_f1_precision_mean", "Precision at validation max-F1 threshold, mean"),
        ("valid_max_f1_recall_mean", "Recall at validation max-F1 threshold, mean"),
        ("valid_max_f1_f1_mean", "F1 at validation max-F1 threshold, mean")])
    numeric_sheet(wb, "Genome-wide (numbers)", genome, [
        ("method", "Method"), ("genome", "Genome"), ("model", "Model"), ("calls_by", "Calls by"), ("calls", "Calls"),
        ("calls_per_mb", "Calls per Mb"), ("precision_unambiguous_calls", "Precision of calls"),
        ("ap_unambiguous_candidates", "AP, unambiguous candidates"), ("ap_training_family_loci", "AP, training-family loci"),
        ("ap_heldout_family_loci", "AP, held-out-family loci"), ("positives_training_family", "Positive candidates, training families"),
        ("positives_heldout_family", "Positive candidates, held-out families"), ("all_loci_recall", "Recall, all loci"),
        ("unmasked_loci_recall", "Recall, unmasked loci"), ("unmasked_training_family_loci_recall", "Recall, training families"),
        ("unmasked_heldout_family_loci_recall", "Recall, held-out families"), ("all_loci", "Loci, all"),
        ("unmasked_loci", "Loci, unmasked"), ("unmasked_training_family_loci", "Loci, training families"),
        ("unmasked_heldout_family_loci", "Loci, held-out families")])
    validation = validation_rows(results)
    numeric_sheet(wb, "Validation AP", validation, auto_columns(validation))
    replicates = replicate_rows(results)
    numeric_sheet(wb, "Per model", [r for r in replicates if r["species"] == ""],
                  [c for c in auto_columns(replicates) if c[0] != "species"])
    per_species = species_rows(results)
    numeric_sheet(wb, "Per species (held-out)", per_species, auto_columns(per_species))
    pairs = paired_rows(results)
    numeric_sheet(wb, "Paired differences", pairs, auto_columns(pairs))
    genome_all = genome_all_rows(results)
    numeric_sheet(wb, "Genome-wide (all metrics)", [r for r in genome_all if r["category"] != "reference, not comparable"],
                  auto_columns(genome_all))
    notes_sheet(wb)
    write_csv(out / "all_results.tsv", long_rows([
        ("benchmark_tests", [r for r in summary_rows(results) if r["species"] == ""]),
        ("benchmark_per_species", [r for r in summary_rows(results) if r["species"] != ""]),
        ("benchmark_per_replicate", replicates),
        ("benchmark_validation", validation),
        ("benchmark_paired_differences", pairs),
        ("genome", genome_all),
    ]), delimiter="\t")
    for ws in wb.worksheets:  # print each sheet one page wide, landscape
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    wb.save(args.xlsx)
    pending = sorted({f"{r['tool']} {r['genome']}" for r in genome if r["calls"] == "pending"})
    print(f"wrote {out}/benchmark_tests.csv, genome_wide.csv, benchmark_tests_all_metrics.tsv, genome_wide_all_metrics.tsv, all_results.tsv and {args.xlsx}; pending genome results: {', '.join(pending) or 'none'}")


if __name__ == "__main__":
    main()
